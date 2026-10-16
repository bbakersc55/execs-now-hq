"""The strategy session as an editable Word file (2026-10-15).

Beside the PDF: the same sections and words, from the same filtered content,
so nothing private reaches the Word file that would not reach the PDF. Who may
download it is who may download the PDF: the practice's staff on a session
they may open; no client user, and nothing across practices.
"""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document

from apps.strategy import docx_export, pdf as pdf_service
from apps.strategy.models import StrategyMapRow

from . import registry_config  # noqa: F401
from .conftest import _member
from .factories import ClientCompanyFactory
from .test_module4_acceptance import (  # noqa: F401
    _a_full_session, answer, company, prospect, session, template,
)


def text_of(content: bytes) -> str:
    """Every word in the file: body, tables and footer."""
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts += [cell.text for cell in row.cells]
    parts += [p.text for p in doc.sections[0].footer.paragraphs]
    return "\n".join(parts)


@pytest.mark.django_db
def test_it_carries_the_pdfs_sections_and_words_whole(session):
    full = _a_full_session(session)
    content = docx_export.render_docx(full)
    assert zipfile.is_zipfile(io.BytesIO(content))           # a real .docx
    text = text_of(content)
    # The two-page design's sections, in its order (2026-10-15).
    headings = ("Strategy session", "The mirror",
                "The six components: where the company stands",
                "The next ninety days: the map", "Two paths: how to get there",
                "Before we meet again: your next moves")
    places = [text.index(heading) for heading in headings]
    assert places == sorted(places), "the PDF's order"
    assert "Acme Facilities · Prepared for Dana Reyes" in text
    assert "Solid processes and accountability mechanisms." in text      # the mirror
    assert "To grow, but do it more sustainably." in text
    # The same words as the PDF, from the same place.
    doc_words = pdf_service.context_for(full)["doc"]
    for component in doc_words["components"]:
        assert component["name"] in text
    for box in doc_words["boxes"]:
        assert f"Day {box['days']} · {box['date']}" in text
        for bullet in box["bullets"]:
            assert bullet["text"] in text
    for move in doc_words["next_moves"]:
        assert move["lead"] in text
    assert doc_words["closing"] in text
    # Headings are Word headings, and the tiles, components and map are tables.
    doc = Document(io.BytesIO(content))
    assert any(p.style.name.startswith("Heading") for p in doc.paragraphs)
    assert len(doc.tables) >= 3


@pytest.mark.django_db
def test_nothing_private_reaches_it_that_the_pdf_leaves_out(session, ff, api):
    answer(session, "s1_revenue", {"text": "4.2m"}, note="MARKERNOTE")
    answer(session, "s4_done_right", {"said": "Spot checks."}, note="MARKERDIAGNOSTIC")
    answer(session, "s3_alignment_observation", {"text": "MARKERALIGNMENT"})
    answer(session, "s9_investment_range", {"agreed": True, "notes": "MARKERINVESTMENT"})
    StrategyMapRow.objects.create(
        tenant=session.tenant, session=session, bottleneck="Supervisor overload",
        state=StrategyMapRow.State.ACCEPTED, mechanics_note="MARKERMECHANICS",
        horizon=60, measurable="Inspections per site per month")
    text = text_of(docx_export.render_docx(session))
    for marker in ("MARKERNOTE", "MARKERDIAGNOSTIC", "MARKERALIGNMENT",
                   "MARKERINVESTMENT", "MARKERMECHANICS"):
        assert marker not in text, marker
    assert "Supervisor overload" in text
    # A flag turned on lets exactly that one through, as in the PDF.
    api.as_(ff).patch(f"/api/strategy-sessions/{session.pk}/pdf-flags/",
                      {"mechanics": True}, content_type="application/json")
    session.refresh_from_db()
    text = text_of(docx_export.render_docx(session))
    assert "MARKERMECHANICS" in text and "MARKERNOTE" not in text
    assert "MARKERMECHANICS" in pdf_service.render_html(session)


@pytest.mark.django_db
def test_it_wears_the_practices_branding(session):
    tenant = session.tenant
    tenant.email_display_name = "Acme Advisory"
    tenant.email_header_color = "#0A3A65"
    tenant.brand_footer_text = "12 Main St\nacme.example"
    tenant.save()
    session.refresh_from_db()
    content = docx_export.render_docx(_a_full_session(session))
    doc = Document(io.BytesIO(content))
    text = text_of(content)
    assert "Acme Advisory" in text and "12 Main St" in text and "acme.example" in text
    assert str(doc.styles["Heading 1"].font.color.rgb) == "0A3A65"


@pytest.mark.django_db
def test_the_practice_downloads_it_beside_the_pdf(session, ff, va, api):
    for who in (ff, va):         # the PDF's own rule: generating is not sending
        response = api.as_(who).get(f"/api/strategy-sessions/{session.pk}/docx/")
        assert response.status_code == 200
        assert response["Content-Type"] == docx_export.CONTENT_TYPE
        assert response["Content-Disposition"].startswith(
            'attachment; filename="strategy-session-acme-facilities-')
        assert response["Content-Disposition"].endswith('.docx"')


@pytest.mark.django_db
@pytest.mark.parametrize("role", ["FCC", "ECC"])
def test_no_client_user_reaches_it(session, api, role):
    client_user = _member(session.tenant, role,
                          company=ClientCompanyFactory(tenant=session.tenant))
    response = api.as_(client_user).get(f"/api/strategy-sessions/{session.pk}/docx/")
    assert response.status_code in (403, 404)


@pytest.mark.django_db
def test_nothing_crosses_practices(session, api, tenant_b):
    other = _member(tenant_b, "FF")
    response = api.as_(other).get(f"/api/strategy-sessions/{session.pk}/docx/")
    assert response.status_code == 404
