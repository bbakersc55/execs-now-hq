"""Reading a OneDrive folder of meeting notes, for a practice on Microsoft 365
(2026-10-27; docs/m365_transport.md §6, D2).

The counterpart of `drive.DriveClient`, with **the same methods**, so the
poller (`ingest.poll`), the folder import (`backfill`), parsing and the queue
are Google's code unchanged: `ingest.client_for` hands a Microsoft practice
this client instead. Everything that decides anything stays in `ingest`.

Microsoft Graph, through the practice owner's Microsoft 365 connection with
`Files.Read`. Three things differ from Drive and are handled here:

1. **Delta only from the root.** OneDrive for Business keeps a change feed
   (`delta`) for the whole drive, not for a folder, so the cursor is the
   root's delta link and each changed file is placed by its parent, kept when
   that is the watched folder or one level below it, as on Google.
2. **Notes written during the meeting.** A Teams or Loop notes page is edited
   while people talk. A file changed in the last `SETTLE_MINUTES` holds the
   cursor where it is (the page comes back with no token, so `ingest` keeps
   its place), and the next poll sees it again, finished.
3. **A file edited after it was read** comes round again with a new version.
   It is recorded as skipped, with the reason, rather than read into a second
   proposal for the same meeting.

What only a real tenant can confirm is in docs/m365_transport.md §10.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timedelta, timezone as dt_timezone

from apps.meetings.drive import DriveFile, DrivePage, DriveUnavailable, FolderInfo, SubFolder

#: A file changed more recently than this is still being written.
SETTLE_MINUTES = 20
MAX_PAGES = 10

#: What is read, by extension. OneDrive's `file.mimeType` is reported too, but
#: a Loop page has none worth trusting, so the name decides.
READ_AS = {
    ".docx": "docx",     # Word; Copilot's notes saved as a document
    ".txt": "text",
    ".md": "text",
    ".vtt": "vtt",       # a Teams transcript downloaded or saved
    ".loop": "loop",     # Teams meeting notes, and Copilot's, in Loop
    ".fluid": "loop",    # Loop's older extension
}
SKIP_BY_EXTENSION = {
    ".pdf": "A PDF is not read yet. Save the notes as a Word document or text.",
    ".mp4": "Recordings are not read from OneDrive.",
    ".m4a": "Recordings are not read from OneDrive.",
    ".mp3": "Recordings are not read from OneDrive.",
    ".pptx": "A presentation is not meeting notes.",
    ".xlsx": "A spreadsheet is not meeting notes.",
}


def extension(name: str) -> str:
    match = re.search(r"\.[A-Za-z0-9]+$", name or "")
    return match.group(0).lower() if match else ""


def skip_reason(name: str) -> str:
    ext = extension(name)
    if ext in READ_AS:
        return ""
    return SKIP_BY_EXTENSION.get(ext, f"{ext or 'A file with no extension'} is not read yet.")


def _when(value: str):
    try:
        return datetime.fromisoformat((value or "").replace("Z", "+00:00"))
    except ValueError:
        return None


def docx_text(content: bytes) -> str:
    """A Word document's paragraphs and table cells, in order."""
    from docx import Document

    try:
        document = Document(io.BytesIO(content))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise DriveUnavailable(f"That Word document could not be opened: {exc}") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())


_CUE_TIME = re.compile(r"^\d{1,2}:\d{2}(:\d{2})?[.,]\d{3}\s+-->\s+")
_VOICE = re.compile(r"<v\s+([^>]+)>(.*?)(?:</v>)?$")


def vtt_text(text: str) -> str:
    """A WebVTT transcript as "Speaker: words" lines: cue numbers, times and
    the header dropped, consecutive lines from one speaker joined."""
    out: list[list[str]] = []
    for raw in text.splitlines():
        line = raw.strip()
        if (not line or line == "WEBVTT" or line.startswith("NOTE")
                or _CUE_TIME.match(line) or re.fullmatch(r"[0-9a-f-]{8,}(/\d+-\d+)?|\d+", line)):
            continue
        voice = _VOICE.match(line)
        speaker, words = (voice.group(1).strip(), voice.group(2).strip()) if voice else ("", line)
        words = re.sub(r"<[^>]+>", "", words).strip()
        if not words:
            continue
        if out and out[-1][0] == speaker:
            out[-1][1] += " " + words
        else:
            out.append([speaker, words])
    return "\n".join(f"{s}: {w}" if s else w for s, w in out)


def html_text(html: str) -> str:
    """A Loop page converted to HTML by Graph, as plain text with its line
    breaks: headings, paragraphs, list items and table rows each on a line."""
    from lxml import html as lxml_html

    marked = re.sub(r"(?i)<br\s*/?>|</(p|div|li|h[1-6]|tr)>", "\n", html or "")
    marked = re.sub(r"(?i)<li[^>]*>", "\n- ", marked)
    marked = re.sub(r"(?i)</t[dh]>", " | ", marked)
    try:
        text = lxml_html.fromstring(f"<div>{marked}</div>").text_content()
    except Exception:                                   # pragma: no cover - malformed
        text = re.sub(r"<[^>]+>", " ", marked)
    lines = [re.sub(r"[ \t ]+", " ", line).strip(" |") for line in text.splitlines()]
    return "\n".join(line for line in lines if line.strip())


class OneDriveClient:
    """The real one, on the practice owner's Microsoft 365 connection."""

    def __init__(self, connection, *, now=None):
        self.connection = connection
        self._token = None
        self._folders: dict = {}
        self._now = now

    # --------------------------------------------------------------- calls

    def token(self):
        from apps.crm.services import graph

        if self._token is None:
            try:
                self._token = graph.access_token_for(self.connection)
            except graph.TransportUnavailable as exc:
                raise DriveUnavailable(str(exc)) from exc
        return self._token

    def _get(self, url: str, params=None, *, raw=False, timeout=30):
        from apps.crm.services import graph

        if not url.startswith("http"):
            url = f"{graph.GRAPH}{url}"
        response = graph._call("GET", url, token=self.token(), params=params, timeout=timeout)
        if response.status_code == 404:
            raise DriveUnavailable("OneDrive has no such folder or file for this account.")
        if not response.ok:
            raise DriveUnavailable(
                f"OneDrive refused ({response.status_code}): {graph._said(response)[:200]}")
        return response.content if raw else response.json()

    def now(self):
        return self._now or datetime.now(dt_timezone.utc)

    # ------------------------------------------------------------ folders

    def folder_meta(self, folder_id: str):
        if folder_id not in self._folders:
            try:
                item = self._get(f"/me/drive/items/{folder_id}",
                                 {"$select": "id,name,parentReference"})
                parent = (item.get("parentReference") or {}).get("id")
                self._folders[folder_id] = (item.get("name", ""), [parent] if parent else [])
            except DriveUnavailable:
                self._folders[folder_id] = None
        return self._folders[folder_id]

    def open_folder(self, folder_id: str) -> str:
        item = self._get(f"/me/drive/items/{folder_id}", {"$select": "id,name,folder,deleted"})
        if item.get("deleted"):
            raise DriveUnavailable("That folder has been deleted.")
        if "folder" not in item:
            raise DriveUnavailable(f"“{item.get('name') or folder_id}” is a file, "
                                   "not a folder.")
        return item.get("name", "")

    def children(self, folder_id: str = "") -> dict:
        """One folder and the folders in it, for the picker: its own name and
        path, and each child folder with how many items it holds."""
        path = f"/me/drive/items/{folder_id}" if folder_id else "/me/drive/root"
        item = self._get(path, {"$select": "id,name,folder,parentReference,root"})
        if "folder" not in item:
            raise DriveUnavailable("That is a file, not a folder.")
        rows, url, params = [], f"{path}/children", {
            "$select": "id,name,folder,file,createdDateTime", "$top": "200"}
        for _ in range(MAX_PAGES):
            page = self._get(url, params)
            rows += page.get("value") or []
            url, params = page.get("@odata.nextLink"), None
            if not url:
                break
        parent = item.get("parentReference") or {}
        return {
            "id": item.get("id", ""),
            "name": "OneDrive" if "root" in item else item.get("name", ""),
            "is_root": "root" in item,
            # "/drive/root:/Documents/Meetings" → ["Documents", "Meetings"]
            "path": [p for p in (parent.get("path") or "").split(":", 1)[-1].split("/") if p]
                    + ([] if "root" in item else [item.get("name", "")]),
            "parent_id": "" if "root" in item else parent.get("id", ""),
            "folders": sorted(({"id": r["id"], "name": r.get("name", ""),
                                "items": (r.get("folder") or {}).get("childCount", 0)}
                               for r in rows if "folder" in r), key=lambda f: f["name"].lower()),
            "readable": sum(1 for r in rows if "file" in r and not skip_reason(r.get("name"))),
            "files": sum(1 for r in rows if "file" in r),
        }

    def _count(self, folder_id: str) -> dict:
        files = readable = 0
        folders, oldest, newest, truncated = [], "", "", False
        url, params = f"/me/drive/items/{folder_id}/children", {
            "$select": "id,name,folder,file,createdDateTime", "$top": "200"}
        for _ in range(MAX_PAGES):
            page = self._get(url, params)
            for row in page.get("value") or []:
                if "folder" in row:
                    folders.append(row)
                    continue
                files += 1
                if not skip_reason(row.get("name")):
                    readable += 1
                    created = row.get("createdDateTime", "")
                    oldest = min(oldest or created, created) if created else oldest
                    newest = max(newest, created)
            url, params = page.get("@odata.nextLink"), None
            if not url:
                break
        else:
            truncated = True
        return {"files": files, "readable": readable, "folders": folders,
                "oldest": oldest, "newest": newest, "truncated": truncated}

    def describe_folder(self, folder_id: str) -> FolderInfo:
        """The same report as Drive's: the folder, one level below, counts."""
        name = self.open_folder(folder_id)
        here = self._count(folder_id)
        subfolders, dates = [], [here["oldest"], here["newest"]]
        for child in here.pop("folders"):
            below = self._count(child["id"])
            subfolders.append(SubFolder(folder_id=child["id"], name=child.get("name", ""),
                                        files=below["files"], readable=below["readable"]))
            dates += [below["oldest"], below["newest"]]
        dates = sorted(d for d in dates if d)
        return FolderInfo(folder_id=folder_id, name=name, files=here["files"],
                          readable=here["readable"], truncated=here["truncated"],
                          subfolders=subfolders,
                          oldest=dates[0] if dates else "", newest=dates[-1] if dates else "")

    # ---------------------------------------------------------- the feed

    def start_token(self) -> str:
        """Where to start: now. `token=latest` returns no backlog, only the
        link to ask from (the past is the folder import's)."""
        data = self._get("/me/drive/root/delta", {"token": "latest"})
        link = data.get("@odata.deltaLink", "")
        if not link:
            raise DriveUnavailable("OneDrive gave no place to start reading from.")
        return link

    def _as_file(self, item: dict) -> DriveFile:
        from apps.meetings.models import MeetingSourceFile

        name = item.get("name", "")
        version = str(item.get("cTag") or item.get("eTag") or "")
        reason = skip_reason(name)
        if not reason and version and MeetingSourceFile.objects.filter(
                provider="microsoft", drive_file_id=item.get("id", "")).exclude(
                drive_version=version).exists():
            reason = ("Changed after it was read. The meeting's proposal is from the earlier "
                      "version; read it again from the queue to use this one.")
        creator = ((item.get("createdBy") or {}).get("user") or {})
        parent = (item.get("parentReference") or {}).get("id")
        return DriveFile(
            file_id=item.get("id", ""), version=version, name=name,
            mime_type=(item.get("file") or {}).get("mimeType", ""),
            owner_email=creator.get("email", ""), web_view_link=item.get("webUrl", ""),
            trashed=bool(item.get("deleted")),
            created_time=item.get("createdDateTime", ""),
            parents=[parent] if parent else [], provider="microsoft", skip=reason)

    def changes(self, page_token: str, *, folder_ids, deep_roots=()) -> DrivePage:
        """One page of the drive's changes, kept to the watched folders.

        When a kept file is still being written, the page comes back with the
        files before it and **no token**, so `ingest` keeps the cursor where
        it was and reads this page again next time."""
        wanted = set(folder_ids)
        data = self._get(page_token)
        settled_before = self.now() - timedelta(minutes=SETTLE_MINUTES)
        files = []
        for item in data.get("value") or []:
            if "file" not in item or item.get("deleted"):
                continue
            if (item.get("parentReference") or {}).get("id") not in wanted:
                continue
            changed = _when(item.get("lastModifiedDateTime", ""))
            if changed is not None and changed > settled_before:
                return DrivePage(files=files)
            files.append(self._as_file(item))
        return DrivePage(files=files, next_page_token=data.get("@odata.nextLink", ""),
                         new_start_page_token=data.get("@odata.deltaLink", ""))

    # ---------------------------------------------------------- backfill

    def files_in(self, folder_ids, *, created_from: str = "", page_size: int = 50,
                 max_pages: int = 1) -> list[DriveFile]:
        """Readable files in these folders, oldest first. Graph cannot order a
        folder's children by date across folders, so each is listed and the
        whole sorted here."""
        found = []
        for folder_id in folder_ids:
            url, params = f"/me/drive/items/{folder_id}/children", {
                "$select": ("id,name,file,folder,cTag,eTag,webUrl,createdDateTime,"
                            "lastModifiedDateTime,createdBy,parentReference"),
                "$top": "200"}
            for _ in range(MAX_PAGES):
                page = self._get(url, params)
                for item in page.get("value") or []:
                    if "file" not in item:
                        continue
                    if created_from and (item.get("createdDateTime") or "") < created_from:
                        continue
                    found.append(self._as_file(item))
                url, params = page.get("@odata.nextLink"), None
                if not url:
                    break
        found.sort(key=lambda f: f.created_time or "")
        return found[:page_size * max_pages]

    def search(self, name_pattern: str, *, roots, created_from: str = "",
               page_size: int = 50, max_pages: int = 10) -> list[DriveFile]:
        """Any-depth folders are Google Meet's; OneDrive watches have none."""
        return []

    # ---------------------------------------------------------- one file

    def file(self, file_id: str) -> DriveFile | None:
        try:
            item = self._get(f"/me/drive/items/{file_id}")
        except DriveUnavailable:
            return None
        return self._as_file(item)

    def text_of(self, drive_file: DriveFile) -> str:
        """The document as text: Word through python-docx, a Loop page as the
        HTML Graph converts it to, a transcript without its timings."""
        kind = READ_AS.get(extension(drive_file.name))
        base = f"/me/drive/items/{drive_file.file_id}/content"
        if kind == "loop":
            return html_text(self._get(base, {"format": "html"}, raw=True, timeout=60)
                             .decode("utf-8", "replace"))
        content = self._get(base, raw=True, timeout=60)
        if kind == "docx":
            return docx_text(content)
        text = content.decode("utf-8-sig", "replace")
        return vtt_text(text) if kind == "vtt" else text

